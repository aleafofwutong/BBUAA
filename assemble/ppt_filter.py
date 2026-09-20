"""Filter near-identical consecutive screenshot slides from PPT exports.

Run ``python -m assemble.ppt_filter input.pptx output.pptx`` to clean an
existing export. Only full-page, single-picture slides are candidates;
editable slides and other layouts are always retained.
"""

from __future__ import annotations

import argparse
import io
from pathlib import Path
from typing import Callable, Iterable

from PIL import Image, ImageChops, ImageDraw, ImageOps, ImageStat

DEFAULT_MAX_CHANGED_RATIO = 0.003
DEFAULT_PIXEL_DELTA = 25
DEFAULT_MAX_MEAN_DELTA = 2.0
THUMBNAIL_SIZE = (640, 360)
RESAMPLE_BILINEAR = getattr(Image, "Resampling", Image).BILINEAR
QR_MAX_CHANGED_RATIO = 0.012
QR_MAX_MEAN_DELTA = 3.0


def _thumbnail(data: bytes) -> tuple[tuple[int, int], Image.Image]:
    with Image.open(io.BytesIO(data)) as source:
        image = ImageOps.exif_transpose(source)
        size = image.size
        return size, image.convert("RGB").resize(THUMBNAIL_SIZE, RESAMPLE_BILINEAR)


def _qr_bounds(image: Image.Image) -> tuple[int, int, int, int] | None:
    import cv2
    import numpy as np

    found, corners = cv2.QRCodeDetector().detect(np.asarray(ImageOps.grayscale(image)))
    if not found or corners is None:
        return None
    x1, y1 = corners[0].min(axis=0)
    x2, y2 = corners[0].max(axis=0)
    width, height = x2 - x1, y2 - y1
    # Exclude small QR codes (e.g. a link on an otherwise changing lecture slide).
    if not (image.height * 0.25 <= height <= image.height * 0.8):
        return None
    if not (0.85 <= width / height <= 1.15):
        return None
    if width * height > image.width * image.height * 0.3:
        return None
    return round(x1), round(y1), round(x2), round(y2)


def _similar(
    previous: tuple[tuple[int, int], Image.Image],
    current: tuple[tuple[int, int], Image.Image],
    *,
    max_changed_ratio: float,
    pixel_delta: int,
    max_mean_delta: float,
) -> bool:
    if previous[0] != current[0]:
        return False
    difference = ImageChops.difference(previous[1], current[1])
    channels = difference.split()
    strongest = ImageChops.lighter(channels[0], ImageChops.lighter(channels[1], channels[2]))
    changed = strongest.point(lambda value: 255 if value > pixel_delta else 0)
    changed_ratio = ImageStat.Stat(changed).mean[0] / 255
    mean_delta = sum(ImageStat.Stat(difference).mean) / 3
    if changed_ratio <= max_changed_ratio and mean_delta <= max_mean_delta:
        return True

    first_qr = _qr_bounds(previous[1])
    second_qr = _qr_bounds(current[1])
    if first_qr is None or second_qr is None:
        return False
    if any(abs(a - b) > 12 for a, b in zip(first_qr, second_qr)):
        return False

    mask = Image.new("L", THUMBNAIL_SIZE, 255)
    draw = ImageDraw.Draw(mask)
    for x1, y1, x2, y2 in (first_qr, second_qr):
        draw.rectangle((x1 - 14, y1 - 14, x2 + 14, y2 + 14), fill=0)
    outside_changed = ImageStat.Stat(changed, mask).mean[0] / 255
    outside_delta = sum(ImageStat.Stat(difference, mask).mean) / 3
    return (
        outside_changed <= min(QR_MAX_CHANGED_RATIO, max_changed_ratio * 4)
        and outside_delta <= min(QR_MAX_MEAN_DELTA, max_mean_delta * 1.5)
    )


def filter_adjacent_images(
    images: Iterable[tuple[int, bytes]],
    *,
    max_changed_ratio: float = DEFAULT_MAX_CHANGED_RATIO,
    pixel_delta: int = DEFAULT_PIXEL_DELTA,
    max_mean_delta: float = DEFAULT_MAX_MEAN_DELTA,
) -> tuple[list[tuple[int, bytes]], list[int]]:
    """Keep the first frame of each similar run, comparing with the last kept frame."""
    if not 0 <= max_changed_ratio <= 1:
        raise ValueError("max_changed_ratio must be between 0 and 1")
    if not 0 <= pixel_delta <= 255:
        raise ValueError("pixel_delta must be between 0 and 255")
    if max_mean_delta < 0:
        raise ValueError("max_mean_delta must be non-negative")

    kept: list[tuple[int, bytes]] = []
    removed: list[int] = []
    previous: tuple[tuple[int, int], Image.Image] | None = None
    previous_index: int | None = None
    for index, data in images:
        current = _thumbnail(data)
        is_adjacent = previous_index is not None and index == previous_index + 1
        previous_index = index
        if previous is not None and is_adjacent and _similar(
            previous,
            current,
            max_changed_ratio=max_changed_ratio,
            pixel_delta=pixel_delta,
            max_mean_delta=max_mean_delta,
        ):
            removed.append(index)
            continue
        kept.append((index, data))
        previous = current
    return kept, removed


def filter_timeline_images(
    timeline: list[dict], fetch_image: Callable[[str], bytes],
) -> tuple[list[dict], list[int]]:
    """Filter a timeline without collapsing across images that could not be read."""
    kept: list[dict] = []
    removed: list[int] = []
    previous: tuple[tuple[int, int], Image.Image] | None = None
    for index, slide in enumerate(timeline):
        try:
            current = _thumbnail(fetch_image(slide["img_url"]))
        except (KeyError, OSError, ValueError):
            previous = None
            kept.append(slide)
            continue
        if previous is not None and _similar(
            previous,
            current,
            max_changed_ratio=DEFAULT_MAX_CHANGED_RATIO,
            pixel_delta=DEFAULT_PIXEL_DELTA,
            max_mean_delta=DEFAULT_MAX_MEAN_DELTA,
        ):
            removed.append(index + 1)
            continue
        kept.append(slide)
        previous = current
    return kept, removed


def filter_pptx(source: Path, output: Path, *, max_changed_ratio: float = DEFAULT_MAX_CHANGED_RATIO) -> list[int]:
    """Write a copy of a screenshot PPTX with near-identical adjacent pages removed."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if not 0 <= max_changed_ratio <= 1:
        raise ValueError("max_changed_ratio must be between 0 and 1")
    if source.resolve() == output.resolve():
        raise ValueError("input and output must be different files")
    presentation = Presentation(source)
    removed: list[int] = []
    previous: tuple[tuple[int, int], Image.Image] | None = None
    previous_geometry: tuple[int, int, int, int] | None = None

    for index, slide in enumerate(presentation.slides):
        if len(slide.shapes) != 1 or slide.shapes[0].shape_type != MSO_SHAPE_TYPE.PICTURE:
            previous = None
            previous_geometry = None
            continue
        shape = slide.shapes[0]
        if (
            abs(shape.left) > presentation.slide_width // 50
            or abs(shape.top) > presentation.slide_height // 50
            or shape.width < presentation.slide_width * 0.98
            or shape.height < presentation.slide_height * 0.98
        ):
            previous = None
            previous_geometry = None
            continue
        geometry = (shape.left, shape.top, shape.width, shape.height)
        try:
            current = _thumbnail(shape.image.blob)
        except (OSError, ValueError):
            previous = None
            previous_geometry = None
            continue

        if previous is not None and previous_geometry == geometry and _similar(
            previous,
            current,
            max_changed_ratio=max_changed_ratio,
            pixel_delta=DEFAULT_PIXEL_DELTA,
            max_mean_delta=DEFAULT_MAX_MEAN_DELTA,
        ):
            removed.append(index + 1)
            continue
        previous = current
        previous_geometry = geometry

    # Remove slide relationships from the copy, leaving the original untouched.
    for page in reversed(removed):
        slide_id = presentation.slides._sldIdLst[page - 1]
        presentation.part.drop_rel(slide_id.rId)
        presentation.slides._sldIdLst.remove(slide_id)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    return removed


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="移除 PPTX 中相邻且几乎相同的截图页")
    parser.add_argument("input", type=Path, help="原始 PPTX")
    parser.add_argument("output", type=Path, help="输出 PPTX（不得覆盖原文件）")
    parser.add_argument(
        "--max-changed-ratio", type=float, default=DEFAULT_MAX_CHANGED_RATIO,
        help="允许变化的像素比例，默认 0.003；调小更保守",
    )
    args = parser.parse_args(argv)
    if not 0 <= args.max_changed_ratio <= 1:
        parser.error("--max-changed-ratio must be between 0 and 1")
    try:
        removed = filter_pptx(args.input, args.output, max_changed_ratio=args.max_changed_ratio)
    except (OSError, ValueError) as exc:
        parser.error(str(exc))
    page_summary = str(removed) if len(removed) <= 30 else f"{removed[:30]} ..."
    print(f"已保存 {args.output}，移除 {len(removed)} 页（原页码：{page_summary}）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
