from __future__ import annotations

import io
import json
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import requests
import cv2
from PIL import Image
from pptx import Presentation

from assemble.courses import ClassroomClient
from assemble.ppt_filter import filter_adjacent_images, filter_timeline_images
from assemble.server import create_app


def image_bytes(color: tuple[int, int, int]) -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (320, 180), color).save(output, format="PNG")
    return output.getvalue()


def qr_image_bytes(message: str, *, lecture_changed: bool = False) -> bytes:
    image = Image.new("RGB", (640, 360), "white")
    qr = Image.fromarray(cv2.QRCodeEncoder_create().encode(message)).convert("RGB")
    nearest = getattr(Image, "Resampling", Image).NEAREST
    image.paste(qr.resize((200, 200), nearest), (100, 80))
    if lecture_changed:
        image.paste("black", (450, 100, 550, 180))
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


class FilterAdjacentImagesTests(unittest.TestCase):
    def test_removes_only_near_identical_adjacent_images(self) -> None:
        black = image_bytes((0, 0, 0))
        white = image_bytes((255, 255, 255))

        kept, removed = filter_adjacent_images([
            (0, black),
            (1, black),
            (2, white),
            (3, white),
        ])

        self.assertEqual([index for index, _ in kept], [0, 2])
        self.assertEqual(removed, [1, 3])

    def test_does_not_compare_across_a_missing_page(self) -> None:
        black = image_bytes((0, 0, 0))

        kept, removed = filter_adjacent_images([(0, black), (2, black)])

        self.assertEqual([index for index, _ in kept], [0, 2])
        self.assertEqual(removed, [])

    def test_rejects_invalid_thresholds(self) -> None:
        with self.assertRaises(ValueError):
            filter_adjacent_images([], max_changed_ratio=1.1)

    def test_timeline_keeps_unreadable_images_and_their_neighbors(self) -> None:
        black = image_bytes((0, 0, 0))
        timeline = [{"img_url": url} for url in ("one", "broken", "three")]

        def fetch(url: str) -> bytes:
            if url == "broken":
                raise OSError("image unavailable")
            return black

        kept, removed = filter_timeline_images(timeline, fetch)
        self.assertEqual(kept, timeline)
        self.assertEqual(removed, [])

    def test_ignores_rotating_qr_but_keeps_changed_lecture_content(self) -> None:
        frames = [
            qr_image_bytes("check-in one"),
            qr_image_bytes("check-in two"),
            qr_image_bytes("check-in three", lecture_changed=True),
        ]
        kept, removed = filter_adjacent_images(enumerate(frames))
        self.assertEqual([index for index, _ in kept], [0, 2])
        self.assertEqual(removed, [1])

    def test_ignores_qr_only_when_both_frames_have_it_at_same_position(self) -> None:
        qr = qr_image_bytes("check-in one")
        blank = image_bytes((255, 255, 255))
        kept, removed = filter_adjacent_images([(0, qr), (1, blank)])
        self.assertEqual(len(kept), 2)
        self.assertEqual(removed, [])


class PptTimelineSourceTests(unittest.TestCase):
    def test_stops_when_upstream_ignores_page_number(self) -> None:
        batch = [
            {"created_sec": i, "content": json.dumps({"pptimgurl": f"https://example.org/{i}.jpg"})}
            for i in range(100)
        ]
        session = Mock()
        session.cookies = requests.cookies.RequestsCookieJar()
        session.get.return_value.json.return_value = {"code": 0, "list": batch}
        client = ClassroomClient(session=session, proxies=None)

        slides = client.get_ppt_timeline("course", "sub", "guid")

        self.assertEqual(len(slides), 100)
        self.assertEqual(session.get.call_count, 2)

    def test_collects_distinct_pages(self) -> None:
        first = [
            {"created_sec": i, "content": json.dumps({"pptimgurl": f"https://example.org/{i}.jpg"})}
            for i in range(100)
        ]
        second = [
            {"created_sec": 100, "content": json.dumps({"pptimgurl": "https://example.org/100.jpg"})}
        ]
        session = Mock()
        session.cookies = requests.cookies.RequestsCookieJar()
        session.get.side_effect = [
            Mock(**{"json.return_value": {"code": 0, "list": first, "total": 101}}),
            Mock(**{"json.return_value": {"code": 0, "list": second, "total": 101}}),
        ]
        client = ClassroomClient(session=session, proxies=None)

        slides = client.get_ppt_timeline("course", "sub", "guid")

        self.assertEqual(len(slides), 101)
        self.assertEqual(session.get.call_count, 2)

    def test_respects_upstream_total_even_if_batch_exceeds_page_size(self) -> None:
        batch = [
            {"created_sec": i, "content": json.dumps({"pptimgurl": f"https://example.org/{i}.jpg"})}
            for i in range(111)
        ]
        session = Mock()
        session.cookies = requests.cookies.RequestsCookieJar()
        session.get.return_value.json.return_value = {"code": 0, "list": batch, "total": 111}
        client = ClassroomClient(session=session, proxies=None)

        slides = client.get_ppt_timeline("course", "sub", "guid")

        self.assertEqual(len(slides), 111)
        self.assertEqual(session.get.call_count, 1)


class PptDownloadTests(unittest.TestCase):
    def setUp(self) -> None:
        self.black = image_bytes((0, 0, 0))
        self.white = image_bytes((255, 255, 255))

    def _get(self, filter_similar: bool):
        images = {
            "https://resource.msa.buaa.edu.cn/0.png": self.black,
            "https://resource.msa.buaa.edu.cn/1.png": self.black,
            "https://resource.msa.buaa.edu.cn/2.png": self.white,
        }

        class FakeResponse:
            status_code = 200

            def __init__(self, content: bytes) -> None:
                self.content = content
                self.headers = {"content-type": "image/png"}
                self.text = ""

        class DownloadSession:
            def __init__(self) -> None:
                self.cookies = requests.cookies.RequestsCookieJar()
                self.headers: dict[str, str] = {}

            def get(self, url: str, **_kwargs):
                return FakeResponse(images[url])

        class ClassroomClient:
            def __init__(self) -> None:
                self.session = requests.Session()
                self.proxies = None

            def get_ppt_timeline(self, _course_id: str, _sub_id: str, _guid: str):
                return [{"img_url": url} for url in images]

        fake_client = ClassroomClient()
        with (
            patch("assemble.server.load_cookies", return_value=requests.Session()),
            patch("assemble.server.is_logged_in", return_value=True),
            patch("assemble.server.ClassroomClient.from_cookies", return_value=fake_client),
            patch("requests.Session", DownloadSession),
        ):
            app = create_app(Path("unused-cookies.json"))
            app.testing = True
            return app.test_client().get(
                "/api/courses/ppt/download",
                query_string={
                    "course_id": "course",
                    "sub_id": "sub",
                    "resource_guid": "guid",
                    "filter_similar": "1" if filter_similar else "0",
                },
            )

    def test_download_filters_similar_pages_and_reports_counts(self) -> None:
        response = self._get(filter_similar=True)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["X-PPT-Original-Slides"], "3")
        self.assertEqual(response.headers["X-PPT-Slides"], "2")
        self.assertEqual(response.headers["X-PPT-Removed-Slides"], "1")
        self.assertIn("ppt_course_sub_filtered.pptx", response.headers["Content-Disposition"])
        self.assertEqual(len(Presentation(io.BytesIO(response.data)).slides), 2)

    def test_download_can_keep_all_pages(self) -> None:
        response = self._get(filter_similar=False)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["X-PPT-Slides"], "3")
        self.assertEqual(response.headers["X-PPT-Removed-Slides"], "0")
        self.assertIn("ppt_course_sub.pptx", response.headers["Content-Disposition"])
        self.assertEqual(len(Presentation(io.BytesIO(response.data)).slides), 3)

    def test_timeline_filters_images_and_can_return_original(self) -> None:
        images = {
            "https://resource.msa.buaa.edu.cn/0.png": self.black,
            "https://resource.msa.buaa.edu.cn/1.png": self.black,
            "https://resource.msa.buaa.edu.cn/2.png": self.white,
        }
        upstream = Mock()
        upstream.session = requests.Session()
        upstream.proxies = None
        upstream.get_ppt_timeline.return_value = [
            {"time_sec": i, "img_url": url} for i, url in enumerate(images)
        ]
        response = Mock()
        response.content = b""

        def fake_get(url, **_kwargs):
            response.content = images[url]
            return response

        with (
            patch("assemble.server.load_cookies", return_value=requests.Session()),
            patch("assemble.server.is_logged_in", return_value=True),
            patch("assemble.server.ClassroomClient.from_cookies", return_value=upstream),
            patch("requests.get", side_effect=fake_get),
        ):
            app = create_app(Path("unused-cookies.json"))
            app.testing = True
            client = app.test_client()
            query = {"course_id": "course", "sub_id": "sub", "resource_guid": "guid"}
            filtered = client.get("/api/courses/ppt", query_string=query).get_json()
            original = client.get(
                "/api/courses/ppt", query_string={**query, "filter_similar": "0"},
            ).get_json()

        self.assertEqual(filtered["original_total"], 3)
        self.assertEqual(filtered["total"], 2)
        self.assertEqual(filtered["removed_pages"], [2])
        self.assertEqual([slide["time_sec"] for slide in filtered["list"]], [0, 2])
        self.assertEqual(original["total"], 3)


if __name__ == "__main__":
    unittest.main()
