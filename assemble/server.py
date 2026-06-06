"""BBUAA step2：课程列表与搜索的本地 Web 服务。"""

from __future__ import annotations

import argparse
import json
import threading
import webbrowser
from dataclasses import asdict
from pathlib import Path

from flask import Flask, jsonify, request, send_from_directory

from assemble.courses import ClassroomClient, SearchFilters
from assemble.sso_login import COOKIE_FILE, is_logged_in, load_cookies, login as sso_login, save_cookies

WEB_DIR = Path(__file__).resolve().parent / "web"
LIVE_CACHE_FILE = Path(__file__).resolve().parent / "live_cache.json"
RECORD_DIR = Path(__file__).resolve().parent / "recordings"
RECORD_META_FILE = Path(__file__).resolve().parent / "recordings.json"


def create_app(cookie_path: Path | None = None) -> Flask:
    app = Flask(__name__, static_folder=str(WEB_DIR), static_url_path="")
    client_holder: dict[str, ClassroomClient | None] = {"client": None}
    cookie_file = cookie_path or COOKIE_FILE

    def get_client() -> ClassroomClient:
        if client_holder["client"] is None:
            session = load_cookies(cookie_file)
            if not is_logged_in(session):
                raise RuntimeError(
                    f"未检测到有效登录 cookie，请先运行: python -m assemble.sso_login "
                    f"(cookie 文件: {cookie_file})"
                )
            client_holder["client"] = ClassroomClient.from_cookies(cookie_file)
        return client_holder["client"]

    @app.get("/")
    def index():
        return send_from_directory(WEB_DIR, "index.html")

    @app.get("/api/status")
    def status():
        try:
            client = get_client()
            return jsonify(
                {
                    "ok": True,
                    "user": client.user_name,
                    "tenant_id": client.tenant_id,
                }
            )
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)}), 401

    @app.get("/api/auth/status")
    def auth_status():
        try:
            session = load_cookies(cookie_file)
            if not is_logged_in(session):
                return jsonify({"ok": True, "logged_in": False, "cookie_file": str(cookie_file)})
            client = ClassroomClient.from_cookies(cookie_file)
            return jsonify(
                {
                    "ok": True,
                    "logged_in": True,
                    "user": client.user_name,
                    "tenant_id": client.tenant_id,
                    "cookie_file": str(cookie_file),
                }
            )
        except FileNotFoundError:
            return jsonify({"ok": True, "logged_in": False, "cookie_file": str(cookie_file)})
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)}), 500

    @app.route("/api/auth/logout", methods=["GET", "POST"])
    def auth_logout():
        cf = cookie_file
        try:
            if cf.exists():
                cf.unlink()
            client_holder["client"] = None
            return jsonify({"ok": True, "msg": "已登出，cookie 文件已删除"})
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)}), 500

    @app.post("/api/auth/login")
    def auth_login():
        payload = request.get_json(silent=True) or {}
        username = str(payload.get("username", "")).strip()
        password = str(payload.get("password", "")).strip()
        if not username or not password:
            return jsonify({"ok": False, "error": "用户名和密码不能为空"}), 400

        try:
            session = sso_login(username, password)
            save_cookies(session, cookie_file)
            client_holder["client"] = None
            client = get_client()
            return jsonify(
                {
                    "ok": True,
                    "logged_in": True,
                    "user": client.user_name,
                    "tenant_id": client.tenant_id,
                    "cookie_file": str(cookie_file),
                }
            )
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)}), 401

    @app.get("/api/meta/terms")
    def meta_terms():
        return jsonify({"list": get_client().list_terms()})

    @app.get("/api/meta/colleges")
    def meta_colleges():
        return jsonify({"list": get_client().list_colleges()})

    @app.get("/api/meta/campuses")
    def meta_campuses():
        return jsonify({"list": get_client().list_campuses()})

    @app.get("/api/meta/buildings")
    def meta_buildings():
        campus_id = request.args.get("campus_id", "").strip()
        if not campus_id:
            return jsonify({"list": []})
        return jsonify({"list": get_client().list_buildings(campus_id)})

    @app.get("/api/meta/rooms")
    def meta_rooms():
        building_id = request.args.get("building_id", "").strip()
        if not building_id:
            return jsonify({"list": []})
        return jsonify({"list": get_client().list_rooms(building_id)})

    @app.get("/api/courses/search")
    def courses_search():
        filters = SearchFilters(
            title=request.args.get("title", "").strip(),
            realname=request.args.get("realname", "").strip(),
            course_code=request.args.get("course_code", "").strip(),
            kkxycode=request.args.get("kkxycode", "").strip() or request.args.get("kkxy_code", "").strip(),
            term=request.args.get("term", "").strip(),
            create_at=request.args.get("create_at", "").strip() or request.args.get("search_time", "").strip(),
            search_time=request.args.get("search_time", "").strip() or request.args.get("create_at", "").strip(),
            campus_id=request.args.get("campus_id", "").strip(),
            building_id=request.args.get("building_id", "").strip(),
            room_id=request.args.get("room_id", "").strip(),
            page=max(1, int(request.args.get("page", 1))),
            per_page=min(50, max(1, int(request.args.get("per_page", 20)))),
        )
        try:
            # 有状态筛选时拉全量（当天课程通常 < 500），避免分页导致漏数据
            status_filter = request.args.get("status_filter", "").strip()
            if status_filter:
                filters.per_page = 500

            result = get_client().search(filters)

            if status_filter and result.get("list"):
                label_map = {"live": "直播中", "playback": "回放", "generating": "回放生成中"}
                target = label_map.get(status_filter, "")
                if target:
                    result["list"] = [c for c in result["list"] if c.get("status_label") == target]
                    result["total"] = len(result["list"])
            return jsonify({**result, "filters": asdict(filters)})
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)}), 500

    @app.get("/api/courses/detail")
    def course_detail():
        course_id = request.args.get("course_id", "").strip()
        sub_id = request.args.get("sub_id", "").strip()
        search_time = request.args.get("search_time", "").strip()
        if not course_id or not sub_id:
            return jsonify({"ok": False, "error": "需要 course_id 和 sub_id"}), 400
        try:
            client = get_client()
            detail = client.get_course_detail(
                course_id, sub_id, search_time=search_time
            )
            if detail is None:
                return jsonify({"ok": False, "error": "未找到该课程详情"}), 404

            # 收集所有候选 PPT resource_guid（前端会逐个尝试）
            ppt_guids: list[str] = []
            if detail.get("ppt_resource_guid"):
                ppt_guids.append(detail["ppt_resource_guid"])
            if detail.get("sub_resource_guid") and detail["sub_resource_guid"] not in ppt_guids:
                ppt_guids.append(detail["sub_resource_guid"])
            if detail.get("ppt_video") and detail["ppt_video"]["resource_guid"] not in ppt_guids:
                ppt_guids.append(detail["ppt_video"]["resource_guid"])
            if detail.get("ppt_segment") and detail["ppt_segment"]["resource_guid"] not in ppt_guids:
                ppt_guids.append(detail["ppt_segment"]["resource_guid"])
            detail["ppt_guids"] = ppt_guids
            # 主 guid 放在第一位
            detail["ppt_resource_guid"] = ppt_guids[0] if ppt_guids else ""

            # 始终从 livingroom 页面补充数据
            # （livingroom 是浏览器实际加载的页面，包含鼠标悬停后可见的视频地址）
            living = client.get_livingroom_video_urls(course_id, sub_id)
            if living.get("source") != "error":
                # 用 livingroom 数据补充空缺
                if living.get("stream_url") and not detail.get("primary_video_url"):
                    detail["primary_video_url"] = living["stream_url"]
                    detail["is_m3u8"] = True
                    detail["has_video"] = True
                if living.get("playback_url"):
                    if not detail.get("playback_url"):
                        detail["playback_url"] = living["playback_url"]
                    if not detail.get("primary_video_url"):
                        detail["primary_video_url"] = living["playback_url"]
                        detail["is_m3u8"] = False
                        detail["has_video"] = True
                if living.get("ppt_video_url") and not (
                    detail.get("ppt_video") and detail["ppt_video"] and detail["ppt_video"]["preview_url"]
                ):
                    detail["ppt_video"] = {
                        "resource_guid": "",
                        "preview_url": living["ppt_video_url"],
                        "duration": 0,
                        "is_m3u8": False,
                    }
                # 把 livingroom 页面发现的所有 video URLs 也带上
                if living.get("all_video_urls"):
                    detail["_livingroom_video_urls"] = living["all_video_urls"]
                # livingroom 发现的 resource_guids 加入候选列表
                for g in living.get("resource_guids", []):
                    if g not in ppt_guids:
                        ppt_guids.append(g)
                detail["ppt_guids"] = ppt_guids
                if ppt_guids:
                    detail["ppt_resource_guid"] = ppt_guids[0]
                detail["_livingroom_scraped"] = True

            # debug: 在响应中暴露视频 URL 解析情况
            detail["_debug"] = {
                "primary": (detail.get("primary_video_url") or "")[:200],
                "is_m3u8": detail.get("is_m3u8", False),
                "trans_socket": (detail.get("trans_socket_url") or "")[:200],
                "playback": (detail.get("playback_url") or "")[:200],
                "has_video": detail.get("has_video", False),
            }
            return jsonify({"ok": True, "detail": detail})
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)}), 500

    @app.get("/api/courses/ppt")
    def course_ppt():
        course_id = request.args.get("course_id", "").strip()
        sub_id = request.args.get("sub_id", "").strip()
        resource_guid = request.args.get("resource_guid", "").strip()
        # 支持逗号分隔的多个 guid，逐个尝试
        alt_guids = [g.strip() for g in request.args.get("alt_guids", "").split(",") if g.strip()]
        if not course_id or not sub_id:
            return jsonify({"ok": False, "error": "需要 course_id, sub_id"}), 400

        client = get_client()
        guids_to_try = [resource_guid] + alt_guids if resource_guid else alt_guids
        guids_to_try = list(dict.fromkeys(g for g in guids_to_try if g))  # 去重

        last_error = ""
        for guid in guids_to_try:
            try:
                timeline = client.get_ppt_timeline(course_id, sub_id, guid)
                if timeline:
                    return jsonify({
                        "ok": True,
                        "total": len(timeline),
                        "list": timeline,
                        "used_guid": guid,
                    })
            except Exception as exc:
                last_error = str(exc)

        if last_error:
            return jsonify({"ok": False, "error": last_error}), 500
        return jsonify({"ok": True, "total": 0, "list": [], "msg": "所有 resource_guid 均无 PPT 数据"})

    @app.get("/api/courses/ppt/download")
    def course_ppt_download():
        """将 PPT 幻灯片打包为 .pptx 文件下载。"""
        import io as _io
        import sys as _sys
        import traceback as _traceback

        def _log(msg: str) -> None:
            print(f"[PPT_DL] {msg}", file=_sys.stderr, flush=True)

        try:
            course_id = request.args.get("course_id", "").strip()
            sub_id = request.args.get("sub_id", "").strip()
            resource_guid = request.args.get("resource_guid", "").strip()
            alt_guids = [g.strip() for g in request.args.get("alt_guids", "").split(",") if g.strip()]
            if not course_id or not sub_id:
                return jsonify({"ok": False, "error": "需要 course_id, sub_id"}), 400

            _log(f"开始下载 PPT: course={course_id} sub={sub_id} guid={resource_guid} alt={alt_guids}")

            client = get_client()
            guids_to_try = [resource_guid] + alt_guids if resource_guid else alt_guids
            guids_to_try = list(dict.fromkeys(g for g in guids_to_try if g))
            _log(f"候选 GUID: {guids_to_try}")

            timeline = []
            for guid in guids_to_try:
                try:
                    timeline = client.get_ppt_timeline(course_id, sub_id, guid)
                    _log(f"GUID {guid}: 获取到 {len(timeline)} 张幻灯片")
                    if timeline:
                        break
                except Exception as exc:
                    _log(f"GUID {guid}: 异常 {exc}")

            if not timeline:
                _log("失败: 无 PPT 时间轴数据")
                return jsonify({"ok": False, "error": "没有可下载的 PPT 数据"}), 404

            _log(f"第一张 URL 示例: {timeline[0].get('img_url', 'N/A')[:120]}")

            # Debug: 检查 session cookies
            cookie_names = [c.name for c in client.session.cookies]
            _log(f"Session cookies ({len(cookie_names)}): {cookie_names}")

            import requests as _requests

            dl_session = _requests.Session()
            for c in client.session.cookies:
                dl_session.cookies.set(
                    c.name, c.value,
                    domain=".msa.buaa.edu.cn",
                    path=c.path or "/",
                )
            dl_session.headers.update({
                "Referer": "https://classroom.msa.buaa.edu.cn/",
                "Origin": "https://classroom.msa.buaa.edu.cn",
                "User-Agent": client.session.headers.get("User-Agent", "Mozilla/5.0"),
            })

            # 测试：先下载第一张图片看是否成功
            test_url = timeline[0].get("img_url", "")
            _log(f"测试下载第1张: {test_url[:120]}")
            try:
                test_resp = dl_session.get(test_url, proxies=client.proxies, timeout=20)
                is_jpg = test_resp.content.startswith(b'\xff\xd8')
                _log(f"测试结果: HTTP {test_resp.status_code}, Content-Type={test_resp.headers.get('content-type','?')}, 大小={len(test_resp.content)} bytes, isJPEG={is_jpg}")
            except Exception as exc:
                _log(f"测试下载异常: {exc}")
                _log(_traceback.format_exc())

            # 逐个下载
            results: list[tuple[int, bytes]] = []
            failed_status: dict[int, int] = {}
            debug_samples: list[dict] = []

            for i, slide in enumerate(timeline):
                img_url = slide.get("img_url", "")
                if not img_url:
                    continue
                try:
                    resp = dl_session.get(img_url, proxies=client.proxies, timeout=20)
                    if resp.status_code == 200:
                        data = resp.content
                        if data.startswith(b'\xff\xd8') or data.startswith(b'\x89PNG') or data.startswith(b'GIF8'):
                            results.append((i, data))
                        else:
                            failed_status[-2] = failed_status.get(-2, 0) + 1
                            if len(debug_samples) < 3:
                                debug_samples.append({"idx": i, "status": resp.status_code, "size": len(data), "magic": data[:4].hex(), "text": data[:200].decode('latin-1')})
                    else:
                        failed_status[resp.status_code] = failed_status.get(resp.status_code, 0) + 1
                        if len(debug_samples) < 3:
                            debug_samples.append({"idx": i, "status": resp.status_code, "size": len(resp.content), "text": resp.text[:200]})
                except Exception as exc:
                    failed_status[-1] = failed_status.get(-1, 0) + 1
                    if len(debug_samples) < 3:
                        debug_samples.append({"idx": i, "error": str(exc)})

            _log(f"下载结果: 成功={len(results)}, 失败统计={failed_status}, 样本={debug_samples}")

            if not results:
                detail = ", ".join(f"HTTP {k}: {v}次" if k > 0 else f"网络错误: {v}次" if k == -1 else f"非图片: {v}次" for k, v in failed_status.items())
                return jsonify({
                    "ok": False,
                    "error": f"所有 {len(timeline)} 张 PPT 图片下载失败",
                    "detail": detail,
                    "debug_samples": debug_samples,
                }), 502

            # ---- PIL 校验 ----
            valid: list[tuple[int, bytes]] = []
            try:
                from PIL import Image as _Image
                _HAS_PIL = True
            except ImportError:
                _HAS_PIL = False

            _log(f"PIL available: {_HAS_PIL}")
            for idx, img_data in results:
                if _HAS_PIL:
                    try:
                        im = _Image.open(_io.BytesIO(img_data))
                        im.load()
                        im = im.convert("RGB")
                        buf_img = _io.BytesIO()
                        im.save(buf_img, format="JPEG", quality=95)
                        valid.append((idx, buf_img.getvalue()))
                    except Exception as exc:
                        _log(f"PIL 转换失败 [{idx}]: {exc}")
                        continue
                else:
                    valid.append((idx, img_data))

            _log(f"PIL 校验后: {len(valid)} 张有效")

            if not valid:
                return jsonify({"ok": False, "error": "所有图片校验失败"}), 502

            # ---- 生成 .pptx ----
            from pptx import Presentation as _Presentation
            from pptx.util import Inches as _Inches

            prs = _Presentation()
            prs.slide_width = _Inches(13.333)
            prs.slide_height = _Inches(7.5)
            blank = prs.slide_layouts[6]
            sw = int(prs.slide_width)
            sh = int(prs.slide_height)

            for _, img_bytes in valid:
                slide = prs.slides.add_slide(blank)
                stream = _io.BytesIO(img_bytes)
                slide.shapes.add_picture(stream, 0, 0, width=sw, height=sh)
                stream.close()

            buf = _io.BytesIO()
            prs.save(buf)
            buf.seek(0)
            _log(f"PPTX 生成成功: {len(valid)} 页, {buf.getbuffer().nbytes} bytes")

            from flask import send_file as _send_file
            return _send_file(
                buf,
                mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                as_attachment=True,
                download_name=f"ppt_{course_id}_{sub_id}.pptx",
            )
        except Exception as exc:
            _log(f"未处理异常: {exc}")
            _log(_traceback.format_exc())
            return jsonify({"ok": False, "error": str(exc), "trace": _traceback.format_exc()[-500:]}), 500

    @app.get("/player/")
    def player_page():
        return send_from_directory(WEB_DIR, "player.html")

    @app.get("/api/proxy/image")
    def proxy_image():
        """代理 PPT 图片请求，携带 classroom 认证 cookie。"""
        from urllib.parse import unquote as _unquote

        url = _unquote(request.args.get("url", ""))
        if not url:
            return jsonify({"ok": False, "error": "缺少 url 参数"}), 400
        allowed = ("resource.msa.buaa.edu.cn", "www.msa.buaa.edu.cn", "video.msa.buaa.edu.cn")
        if not any(host in url for host in allowed):
            return jsonify({"ok": False, "error": "不允许的域名"}), 403
        try:
            client = get_client()
            cookie_header = "; ".join(
                f"{c.name}={c.value}"
                for c in client.session.cookies
            )
            import requests as _requests
            resp = _requests.get(
                url,
                headers={
                    "Referer": "https://classroom.msa.buaa.edu.cn/",
                    "Origin": "https://classroom.msa.buaa.edu.cn",
                    "User-Agent": client.session.headers.get(
                        "User-Agent",
                        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
                    ),
                    "Cookie": cookie_header,
                },
                proxies=client.proxies,
                timeout=30,
                stream=True,
            )
            if resp.status_code >= 400:
                return jsonify({
                    "ok": False,
                    "error": f"上游返回 {resp.status_code}",
                }), 502
            from flask import Response as _Response

            return _Response(
                resp.iter_content(chunk_size=8192),
                content_type=resp.headers.get("content-type", "image/jpeg"),
                headers={
                    "Cache-Control": "public, max-age=3600",
                    "Content-Length": resp.headers.get("content-length", ""),
                },
            )
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)}), 502

    @app.get("/api/proxy/video")
    def proxy_video():
        """代理视频流请求，携带 classroom 认证 cookie。
        解决跨域 cookie 问题：浏览器从 127.0.0.1 加载 video 时
        不会向 resource.msa.buaa.edu.cn 发送认证 cookie。
        支持 Range 请求（视频拖动/seek）。
        """
        from urllib.parse import unquote as _unquote

        url = _unquote(request.args.get("url", ""))
        if not url:
            return jsonify({"ok": False, "error": "缺少 url 参数"}), 400

        allowed = (
            "resource.msa.buaa.edu.cn",
            "lmt",
            "buaa.edu.cn",
            "msa.buaa.edu.cn",
        )
        if not any(host in url for host in allowed):
            return jsonify({"ok": False, "error": "不允许的域名"}), 403

        try:
            # resource.msa.buaa.edu.cn 的视频需要 clientUUID 参数
            if "resource.msa.buaa.edu.cn" in url and "clientUUID=" not in url:
                import uuid as _uuid
                sep = "&" if "?" in url else "?"
                url = f"{url}{sep}clientUUID={_uuid.uuid4()}"

            client = get_client()
            # 关键：requests.Session 的 cookie 是按域名隔离的。
            # classroom.msa.buaa.edu.cn 的 cookie 不会自动发给 resource.msa.buaa.edu.cn。
            # 必须把所有 cookie 拼成 Cookie 头手动注入（模拟浏览器的 same-site 行为）。
            cookie_header = "; ".join(
                f"{c.name}={c.value}"
                for c in client.session.cookies
            )
            proxy_headers = {
                "Referer": "https://classroom.msa.buaa.edu.cn/",
                "Origin": "https://classroom.msa.buaa.edu.cn",
                "User-Agent": client.session.headers.get(
                    "User-Agent",
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
                ),
                "Cookie": cookie_header,
            }
            range_header = request.headers.get("Range", "")
            if range_header:
                proxy_headers["Range"] = range_header

            # 用裸 requests 而非 session，避免 session 的域名级 cookie jar
            # 覆盖我们手动拼的 Cookie 头（resource.msa.buaa.edu.cn 需要全量 cookie）
            import requests as _requests
            resp = _requests.get(
                url,
                headers=proxy_headers,
                proxies=client.proxies,
                timeout=60,
                stream=True,
            )
            if resp.status_code >= 400:
                return jsonify({
                    "ok": False,
                    "error": f"上游返回 {resp.status_code}",
                    "detail": resp.text[:500],
                }), 502

            content_type = resp.headers.get("content-type", "video/mp4")

            # ---- HLS 播放列表：重写分片 URL，全部走代理 ----
            is_m3u8 = "mpegurl" in content_type or url.endswith(".m3u8")
            if is_m3u8:
                from urllib.parse import quote as _url_quote, urljoin as _urljoin
                raw = resp.text
                # 解析原始 m3u8 的 base URL（用于解析相对路径）
                base_url = url if url.endswith("/") else url.rsplit("/", 1)[0] + "/"

                def _make_proxy(seg: str) -> str:
                    if seg.startswith("/api/proxy/"):
                        return seg
                    # 相对路径 → 绝对路径
                    if not seg.startswith("http"):
                        seg = _urljoin(base_url, seg)
                    return "/api/proxy/video?url=" + _url_quote(seg, safe="")

                lines = raw.split("\n")
                rewritten_lines = []
                for line in lines:
                    stripped = line.rstrip("\r")
                    if stripped and not stripped.startswith("#"):
                        # 非注释行 = 分片 URL，改写为代理
                        rewritten_lines.append(_make_proxy(stripped))
                    else:
                        rewritten_lines.append(stripped)
                rewritten = "\n".join(rewritten_lines)

                from flask import Response as _Response
                return _Response(
                    rewritten,
                    status=resp.status_code,
                    headers={
                        "Content-Type": content_type,
                        "Access-Control-Allow-Origin": "*",
                        "Cache-Control": "no-cache",
                    },
                )

            content_length = resp.headers.get("content-length", "")

            from flask import Response as _Response

            def generate():
                for chunk in resp.iter_content(chunk_size=65536):
                    if chunk:
                        yield chunk

            response_headers = {
                "Content-Type": content_type,
                "Accept-Ranges": "bytes",
                "Cache-Control": "public, max-age=3600",
            }
            if content_length:
                response_headers["Content-Length"] = content_length
            if resp.status_code == 206:
                response_headers["Content-Range"] = resp.headers.get(
                    "Content-Range", ""
                )

            return _Response(
                generate(),
                status=resp.status_code,
                headers=response_headers,
                direct_passthrough=True,
            )
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)}), 502

    # ---- 直播缓存 ----
    # 校区与教学楼知识库（含复合楼名如 主-北、主-M）
    _CAMPUSES = [
        ("学院路校区", ["A", "B", "C", "D", "E", "F", "G",
                        "主南", "主北", "主M", "主",
                        "(一)", "(三)", "(五)"]),
        ("沙河校区",   ["J0", "J1", "J2", "J3", "J4", "J5",
                        "SH2", "SH3"]),
        ("杭州校区",   ["教学一号楼", "教学二号楼", "科研一号楼", "科研二号楼"]),
        ("腾讯会议",   []),
    ]

    # 复合楼名：实际格式是 "主-北"、"主-M" 但规范化为 "主北"、"主M"
    _COMPOUND_BUILDINGS = {
        "主-北": "主北", "主-M": "主M",
        "主-南": "主南",
    }

    def _parse_room(room_name: str) -> dict:
        """解析教室名 → {campus, building, room}。"""
        raw = room_name.strip()
        if "腾讯会议" in raw:
            return {"campus": "腾讯会议", "building": "", "room": raw}

        for campus, buildings in _CAMPUSES:
            if campus == "腾讯会议":
                continue
            if raw.startswith(campus):
                rest = raw[len(campus):]
                building = ""
                room = rest

                # 先尝试复合楼名（"主-北202" → building="主北", room="202"）
                for compound, canonical in sorted(_COMPOUND_BUILDINGS.items(), key=lambda x: -len(x[0])):
                    if rest.startswith(compound):
                        building = canonical
                        after = rest[len(compound):]
                        if after and after[0] in "-—－":
                            after = after[1:]
                        room = after if after else rest
                        return {"campus": campus, "building": building, "room": room}

                # 匹配已知教学楼（长名优先）
                for b in sorted(buildings, key=len, reverse=True):
                    if rest.startswith(b):
                        building = b
                        after = rest[len(b):]
                        if after and after[0] in "-—－":
                            after = after[1:]
                        room = after if after else rest
                        break
                return {"campus": campus, "building": building, "room": room}

        # 兜底
        if "-" in raw:
            parts = raw.rsplit("-", 1)
            return {"campus": "", "building": parts[0], "room": parts[1]}
        return {"campus": "", "building": "", "room": raw}

    def _location_sort_key(entry: dict) -> tuple:
        loc = _parse_room(entry.get("room_name", ""))
        return (loc["campus"], loc["building"], loc["room"], entry.get("title", ""))

    def _read_live_cache() -> dict:
        try:
            if LIVE_CACHE_FILE.exists():
                return json.loads(LIVE_CACHE_FILE.read_text(encoding="utf-8"))
        except Exception:
            pass
        return {}

    def _write_live_cache(data: dict) -> None:
        LIVE_CACHE_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    @app.get("/api/live/cache")
    def live_cache_list():
        """获取缓存的直播链接。支持 ?sub_id= 或 ?room= 查询。"""
        sub_id = request.args.get("sub_id", "").strip()
        room = request.args.get("room", "").strip()
        cache = _read_live_cache()
        if sub_id:
            entry = cache.get(sub_id)
            if entry:
                return jsonify({"ok": True, "entry": entry, "matched_by": "sub_id"})
            # sub_id 未命中，尝试用教室名匹配
            for v in cache.values():
                if v.get("room_name", "") == room:
                    return jsonify({"ok": True, "entry": v, "matched_by": "room_name"})
            return jsonify({"ok": False, "error": "未找到"}), 404
        if room:
            for v in cache.values():
                if v.get("room_name", "") == room:
                    return jsonify({"ok": True, "entry": v, "matched_by": "room_name"})
            return jsonify({"ok": False, "error": "未找到"}), 404
        items = list(cache.values())
        items.sort(key=_location_sort_key)
        return jsonify({"ok": True, "list": items})

    @app.post("/api/live/cache")
    def live_cache_save():
        """缓存/更新一个直播链接。相同 sub_id 会覆盖旧链接。"""
        payload = request.get_json(silent=True) or {}
        sub_id = str(payload.get("sub_id", "")).strip()
        if not sub_id:
            return jsonify({"ok": False, "error": "缺少 sub_id"}), 400

        live_url = str(payload.get("live_url", "")).strip()
        if not live_url:
            return jsonify({"ok": False, "error": "缺少 live_url"}), 400

        cache = _read_live_cache()
        old = cache.get(sub_id, {})

        # 验证新链接是否可用
        new_works = False
        try:
            client = get_client()
            import requests as _requests
            test_resp = _requests.get(live_url, headers={
                "Referer": "https://classroom.msa.buaa.edu.cn/",
                "User-Agent": client.session.headers.get("User-Agent", "Mozilla/5.0"),
            }, proxies=client.proxies, timeout=10)
            new_works = test_resp.status_code == 200
        except Exception:
            pass

        room_name = str(payload.get("room_name", old.get("room_name", ""))).strip()
        loc = _parse_room(room_name)
        entry = {
            "course_id": str(payload.get("course_id", old.get("course_id", ""))).strip(),
            "sub_id": sub_id,
            "title": str(payload.get("title", old.get("title", ""))).strip(),
            "lecturer_name": str(payload.get("lecturer_name", old.get("lecturer_name", ""))).strip(),
            "room_name": room_name,
            "campus": loc["campus"],
            "building": loc["building"],
            "room": loc["room"],
            "sub_title": str(payload.get("sub_title", old.get("sub_title", ""))).strip(),
            "live_url": live_url if new_works else old.get("live_url", live_url),
            "cached_at": old.get("cached_at", ""),
            "last_checked": __import__("datetime").datetime.now().isoformat(),
            "url_working": new_works,
            "search_time": str(payload.get("search_time", old.get("search_time", ""))).strip(),
        }
        if not entry["cached_at"]:
            entry["cached_at"] = entry["last_checked"]

        cache[sub_id] = entry
        _write_live_cache(cache)
        return jsonify({"ok": True, "saved": entry, "url_working": new_works})

    @app.delete("/api/live/cache")
    def live_cache_delete():
        sub_id = request.args.get("sub_id", "").strip()
        if not sub_id:
            return jsonify({"ok": False, "error": "缺少 sub_id"}), 400
        cache = _read_live_cache()
        if sub_id in cache:
            del cache[sub_id]
            _write_live_cache(cache)
            return jsonify({"ok": True, "msg": "已删除"})
        return jsonify({"ok": False, "error": "未找到"}), 404

    @app.get("/live-space/")
    def live_space_page():
        return send_from_directory(WEB_DIR, "live_space.html")

    # ---- 直播录制 ----
    import subprocess as _sp
    _active_recordings: dict[str, dict] = {}  # sub_id → {proc, info}

    def _read_record_meta() -> dict:
        try:
            if RECORD_META_FILE.exists():
                return json.loads(RECORD_META_FILE.read_text(encoding="utf-8"))
        except Exception:
            pass
        return {}

    def _write_record_meta(data: dict) -> None:
        RECORD_META_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    @app.post("/api/record/start")
    def record_start():
        payload = request.get_json(silent=True) or {}
        sub_id = str(payload.get("sub_id", "")).strip()
        m3u8_url = str(payload.get("url", "")).strip()
        title = str(payload.get("title", "")).strip()
        if not sub_id or not m3u8_url:
            return jsonify({"ok": False, "error": "缺少 sub_id 或 url"}), 400

        if sub_id in _active_recordings:
            return jsonify({"ok": False, "error": "该课程已在录制中"}), 409

        RECORD_DIR.mkdir(parents=True, exist_ok=True)
        ts = __import__("datetime").datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = RECORD_DIR / f"{sub_id}_{ts}.mp4"

        # ffmpeg 录制 HLS 流
        cmd = [
            "ffmpeg", "-y",
            "-user_agent", "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
            "-headers", f"Referer: https://classroom.msa.buaa.edu.cn/\r\nOrigin: https://classroom.msa.buaa.edu.cn",
            "-i", m3u8_url,
            "-c", "copy",
            "-bsf:a", "aac_adtstoasc",
            str(output_path),
        ]
        try:
            proc = _sp.Popen(cmd, stdout=_sp.DEVNULL, stderr=_sp.DEVNULL)
        except FileNotFoundError:
            return jsonify({"ok": False, "error": "未安装 ffmpeg，请先安装"}), 500

        info = {
            "sub_id": sub_id,
            "title": title,
            "url": m3u8_url,
            "output": str(output_path),
            "started_at": __import__("datetime").datetime.now().isoformat(),
            "status": "recording",
        }
        _active_recordings[sub_id] = {"proc": proc, "info": info}

        # 持久化
        meta = _read_record_meta()
        meta[sub_id] = info
        _write_record_meta(meta)

        return jsonify({"ok": True, "msg": "录制已开始", "info": info})

    @app.post("/api/record/stop")
    def record_stop():
        sub_id = str(request.get_json(silent=True).get("sub_id", "")).strip() if request.get_json(silent=True) else request.args.get("sub_id", "").strip()
        if not sub_id:
            return jsonify({"ok": False, "error": "缺少 sub_id"}), 400

        entry = _active_recordings.pop(sub_id, None)
        if entry:
            proc = entry["proc"]
            proc.terminate()
            try:
                proc.wait(timeout=5)
            except Exception:
                proc.kill()
            entry["info"]["status"] = "stopped"
            entry["info"]["stopped_at"] = __import__("datetime").datetime.now().isoformat()

        meta = _read_record_meta()
        if sub_id in meta:
            meta[sub_id]["status"] = "stopped"
            meta[sub_id]["stopped_at"] = __import__("datetime").datetime.now().isoformat()
            _write_record_meta(meta)

        return jsonify({"ok": True, "msg": "录制已停止", "info": meta.get(sub_id, {})})

    @app.get("/api/record/list")
    def record_list():
        meta = _read_record_meta()
        # 合并活跃录制状态
        for sub_id, entry in _active_recordings.items():
            if sub_id in meta:
                meta[sub_id]["status"] = "recording"
        items = list(meta.values())
        items.sort(key=lambda x: x.get("started_at", ""), reverse=True)
        return jsonify({"ok": True, "list": items})

    @app.delete("/api/record/delete")
    def record_delete():
        sub_id = request.args.get("sub_id", "").strip()
        if not sub_id:
            return jsonify({"ok": False, "error": "缺少 sub_id"}), 400
        # 先停掉正在录制的
        if sub_id in _active_recordings:
            proc = _active_recordings[sub_id]["proc"]
            proc.terminate()
            try: proc.wait(timeout=5)
            except: proc.kill()
            del _active_recordings[sub_id]
        meta = _read_record_meta()
        entry = meta.pop(sub_id, None)
        if entry:
            # 删除录制的文件
            output = entry.get("output", "")
            if output:
                try: Path(output).unlink(missing_ok=True)
                except: pass
            _write_record_meta(meta)
            return jsonify({"ok": True, "msg": "已删除"})
        return jsonify({"ok": False, "error": "未找到"}), 404

    @app.post("/api/shutdown")
    def shutdown():
        shutdown_func = request.environ.get("werkzeug.server.shutdown")
        if shutdown_func is None:
            return jsonify({"ok": False, "error": "当前运行环境不支持关闭回调"}), 500
        threading.Timer(0.2, shutdown_func).start()
        return jsonify({"ok": True, "msg": "服务正在退出"})

    return app


def main() -> None:
    parser = argparse.ArgumentParser(description="启动 BBUAA 课程搜索 Web 界面")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=8765)
    parser.add_argument("--cookie", default=str(COOKIE_FILE))
    parser.add_argument("--no-browser", action="store_true")
    args = parser.parse_args()

    app = create_app(Path(args.cookie))
    url = f"http://{args.host}:{args.port}/"
    print(f"[+] BBUAA 课程浏览器: {url}")
    if not args.no_browser:
        webbrowser.open(url)
    app.run(host=args.host, port=args.port, debug=False)


if __name__ == "__main__":
    main()
